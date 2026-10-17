"""Exporta os checklists de um mês do sistema para o MESMO formato da Base POC 4.0 (abas Operacional e Administrativo,
mesmas colunas e mesma ordem das perguntas), com a coluna final "Validação" (Válido / Inválido / Pendente).

Uso: python exportar_mes.py --ano 2026 --mes 10 --saida "C:\\caminho\\arquivo.xlsx"
Só lê o banco (chave anon do POC_OP/.env); não grava nada nele.
"""
import argparse
import collections
import datetime as dt
import os
import re
import sys
import urllib.parse
import unicodedata

import openpyxl
from openpyxl.styles import Font
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importar_historico as h  # noqa: E402  (catálogos, mapeamento de colunas e acesso REST)

MODELO = h.XLSX
FUSO = dt.timezone(dt.timedelta(hours=-3))  # horário local (Maranhão, sem horário de verão)
SEMANAS = ["1. Primeira", "2. Segunda", "3. Terceira", "4. Quarta"]
VALIDACAO = {"aprovado": "Válido", "reprovado": "Inválido", "pendente": "Pendente"}

# auditagem do sistema -> (aba, catálogo, texto de "Gerência da Equipe" usado na planilha)
TIPOS = {
    "GOMAN": ("Operacional", "goman", "GOMAN"),
    "GSTC": ("Operacional", "gstc", "GSTC"),
    "ADMINISTRATIVO": ("Administrativo", "administrativo", "ADMINISTRATIVO"),
    "ALOJAMENTO": ("Administrativo", "alojamento", "ALOJAMENTO"),
    "LOGISTICA": ("Administrativo", "logistica", "LOGÍSTICA"),
    "OFICINA": ("Administrativo", "oficina", "OFICINA"),
}


def n(s):
    return h.norm(s)


def paginar(path, key, query, ordem):
    out, off = [], 0
    while True:
        q = {**query, "order": ordem, "limit": 1000, "offset": off}
        rows = h.http("GET", path, key, query=q)
        out += rows
        if len(rows) < 1000:
            return out
        off += 1000


def em_lotes(path, key, campo, ids, select, extra=None, ordem="id"):
    out = []
    for i in range(0, len(ids), 100):
        q = {"select": select, campo: "in.(" + ",".join(ids[i:i + 100]) + ")", **(extra or {})}
        out += paginar(path, key, q, ordem)
    return out


def categoria_chave(txt):
    t = n(txt)
    for k, v in (("apr", "apr"), ("regra", "regras"), ("padrinho", "padrinho"), ("epi", "epi"), ("procedimento", "procedimento"),
                 ("veiculo", "veiculos"), ("trabalho em altura", "altura"), ("estrutura", "estruturas"), ("republica", "republicas")):
        if t.startswith(k):
            return v
    return t


def ler_modelo():
    wb = {}
    for aba in ("Operacional", "Administrativo"):
        wb[aba] = pd.read_excel(MODELO, sheet_name=aba).reset_index(drop=True)
    return wb


def blocos_e_constantes(df, valores_ger, i_resumo):
    """Por gerência da planilha: colunas de pergunta usadas e valor fixo (moda != 0) das colunas calculadas."""
    ger = df.iloc[:, 4].astype(str).str.strip()
    blocos, const = {}, {}
    for g in ger.unique():
        sub = df[ger == g]
        usadas = {i for i in range(df.shape[1]) if sub.iloc[:, i].notna().mean() > 0.05}
        blocos[g] = usadas
        c = {}
        for i in range(i_resumo, df.shape[1]):
            v = pd.to_numeric(sub.iloc[:, i], errors="coerce")
            nz = v[v.fillna(0) != 0]
            c[i] = float(nz.mode().iloc[0]) if len(nz) else 0
        const[g] = c
    return blocos, const


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ano", type=int, required=True)
    ap.add_argument("--mes", type=int, required=True)
    ap.add_argument("--saida", required=True)
    args = ap.parse_args()

    key = h.chave_env()
    ini = dt.datetime(args.ano, args.mes, 1, tzinfo=FUSO)
    fim = dt.datetime(args.ano + (args.mes == 12), args.mes % 12 + 1, 1, tzinfo=FUSO)
    ini_utc, fim_utc = ini.astimezone(dt.timezone.utc), fim.astimezone(dt.timezone.utc)

    print("lendo modelo e catálogos…", flush=True)
    catalogos = h.carregar_catalogos()
    modelo = ler_modelo()

    print("lendo checklists do banco…", flush=True)
    subs = paginar("checklist_submissions", key, {
        "select": "id,matricula,observador,auditagem,data,base,equipe,membros,resumo,status",
        "and": f"(data.gte.{ini_utc.strftime('%Y-%m-%dT%H:%M:%SZ')},data.lt.{fim_utc.strftime('%Y-%m-%dT%H:%M:%SZ')})",
    }, "data.asc,id.asc")
    ids = [s["id"] for s in subs]
    resp = em_lotes("checklist_responses", key, "submission_id", ids,
                    "submission_id,pergunta_id,categoria,pergunta,peso,resposta,observacao,resolvido,itens,atribuido_tipo,atribuido_matricula",
                    ordem="submission_id.asc,pergunta_id.asc")
    fotos = em_lotes("checklist_photos", key, "submission_id", ids, "submission_id,tipo,r2_key,sort_order",
                     extra={"tipo": "eq.local"})
    empregados = {e["matricula"]: e for e in paginar("employees", key, {"select": "matricula,coordenador,base"}, "matricula.asc")}
    hier = {}
    for r in paginar("equipes_coordenacao_historico", key, {"select": "prefixo,coordenador,gerencia,base"}, "prefixo.asc"):
        hier[r["prefixo"].upper()] = {"coordenador": r["coordenador"] or "", "gerencia": (r["gerencia"] or "").upper(), "base": (r["base"] or "").upper()}
    for r in paginar("pwa_equipes", key, {"select": "id,prefixo,coordenador,gerencia,base"}, "id.asc"):
        k = (r["prefixo"] or "").strip().upper()
        atual = hier.get(k, {"coordenador": "", "gerencia": "", "base": ""})
        hier[k] = {"coordenador": (r["coordenador"] or "").strip() or atual["coordenador"],
                   "gerencia": (r["gerencia"] or "").strip().upper() or atual["gerencia"],
                   "base": (r["base"] or "").strip().upper() or atual["base"]}
    semanas = {(r["ano"], r["mes"]): (r["fim_s1"], r["fim_s2"], r["fim_s3"])
               for r in paginar("semanas_config", key, {"select": "ano,mes,fim_s1,fim_s2,fim_s3"}, "ano.asc,mes.asc")}
    base_r2 = ""
    for linha in open(h.ENV_FILE, encoding="utf-8"):
        if linha.startswith("VITE_R2_PUBLIC_BASE_URL="):
            base_r2 = linha.split("=", 1)[1].strip().rstrip("/")

    por_sub = collections.defaultdict(list)
    for r in resp:
        por_sub[r["submission_id"]].append(r)
    fotos_sub = collections.defaultdict(list)
    for f in fotos:
        fotos_sub[f["submission_id"]].append(f)
    print(f"{len(subs)} checklists | {len(resp)} respostas | {len(fotos)} fotos", flush=True)

    # ── preparação por aba ──
    prep = {}
    for aba, df in modelo.items():
        cols = list(df.columns)
        i_resumo = next(i for i, c in enumerate(cols) if str(c).strip() == "APR")
        blocos, const = blocos_e_constantes(df, None, i_resumo)
        # valor usado para "Não Conforme" (igual ao do arquivo modelo)
        nc_txt = "Não Conforme"
        for ci in range(13, i_resumo):
            achados = [v for v in df.iloc[:, ci].dropna().unique() if isinstance(v, str) and "onforme" in v and v != "Conforme"]
            if achados:
                nc_txt = achados[0]
                break
        prep[aba] = {"df": df, "cols": cols, "i_resumo": i_resumo, "blocos": blocos, "const": const, "nc": nc_txt, "mapas": {}}

    def mapa_do_tipo(aba, tipo, ger_txt):
        p = prep[aba]
        k = (tipo, ger_txt)
        if k not in p["mapas"]:
            bloco = p["blocos"].get(ger_txt, set())
            m = h.mapa_colunas(tipo, p["df"], catalogos)
            p["mapas"][k] = {i: v for i, v in m.items() if i in bloco}
        return p["mapas"][k]

    linhas = {"Operacional": [], "Administrativo": []}
    # Perguntas do sistema que não existem como coluna no modelo (ex.: as novas de Alojamento): entram depois das colunas
    # originais, antes de "Validação", para não perder respostas.
    extras_texto = {"Operacional": {}, "Administrativo": {}}
    contador = {"Operacional": 0, "Administrativo": 0}
    sem_tipo = collections.Counter()

    for s in subs:
        if s["auditagem"] not in TIPOS:
            sem_tipo[s["auditagem"]] += 1
            continue
        aba, tipo, ger_txt = TIPOS[s["auditagem"]]
        equipe = (s["equipe"] or "").strip()
        hh = hier.get(equipe.upper()) or hier.get((re.split(r"\s+[–-]\s+", equipe.upper())[0]).strip()) or {}
        if aba == "Operacional" and s["auditagem"] == "GSTC" and hh.get("gerencia") == "GERE":
            ger_txt = "GERE"
        p = prep[aba]
        cols, df = p["cols"], p["df"]
        row = [None] * len(cols)
        idx = {n(c): i for i, c in enumerate(cols)}  # (cabeçalhos duplicados: vale o último; usado só p/ metadados)

        def put(nome_inicio, valor):
            for i, c in enumerate(cols):
                if n(c).startswith(nome_inicio):
                    row[i] = valor
                    return

        local = dt.datetime.fromisoformat(s["data"].replace("Z", "+00:00")).astimezone(FUSO).replace(tzinfo=None)
        membros = s["membros"] or []
        rs = por_sub.get(s["id"], [])
        ncs = [r for r in rs if r["resposta"] == "nao_conforme"]

        row[0] = local
        row[1] = int(s["matricula"]) if str(s["matricula"]).isdigit() else s["matricula"]
        row[2] = s["base"]
        row[3] = equipe
        row[4] = ger_txt
        row[5] = len(membros)
        if ncs:
            if any(r["atribuido_tipo"] == "equipe" for r in ncs):
                desvios = len(membros)
            else:
                desvios = len({r["atribuido_matricula"] for r in ncs if r["atribuido_matricula"]})
        else:
            desvios = 0
        row[6] = desvios

        # matrículas e nomes dos colaboradores (só nos blocos da gerência da equipe)
        def colunas_membros(prefixo_n, padrao):
            return [i for i, c in enumerate(cols) if n(c).startswith(prefixo_n) and padrao in n(c)]
        if aba == "Operacional":
            alvo = "goman" if ger_txt == "GOMAN" else "gstc"
            mat_cols = [i for i, c in enumerate(cols) if n(c).startswith("matricula do") and "colaborador" in n(c) and alvo in n(c)]
            nome_cols = [i for i, c in enumerate(cols) if "colaborador" in n(c) and not n(c).startswith("matricula") and i > idx.get("base do obs", 0)
                         and alvo in n(c)]
        else:
            if ger_txt == "LOGÍSTICA":
                mat_cols = [i for i, c in enumerate(cols) if n(c).startswith("matricula do") and "colaborador" in n(c)]
                nome_cols = [i for i, c in enumerate(cols) if "colaborador" in n(c) and not n(c).startswith("matricula") and i > idx.get("base do obs", 0)]
            else:
                mat_cols, nome_cols = [], []
        for pos, m in enumerate(membros):
            if pos < len(mat_cols):
                row[mat_cols[pos]] = int(m["matricula"]) if str(m.get("matricula", "")).isdigit() else m.get("matricula")
            if pos < len(nome_cols):
                row[nome_cols[pos]] = (m.get("nome") or "").upper()

        # respostas nas colunas de pergunta (mesma ordem do modelo)
        mapa = mapa_do_tipo(aba, tipo, ger_txt)
        por_pergunta = {r["pergunta_id"]: r for r in rs}
        usado = set()
        for i, m in sorted(mapa.items()):
            kind, q, item = m
            r = por_pergunta.get(q["id"])
            if r is None:
                continue
            if kind == "ROUO":
                itens = r.get("itens") or []
                if itens and item < len(itens):
                    row[i] = "Conforme" if itens[item]["conforme"] else p["nc"]
                else:
                    row[i] = "Conforme" if r["resposta"] == "conforme" else p["nc"]
            else:
                if q["id"] in usado:
                    continue
                usado.add(q["id"])
                row[i] = "Conforme" if r["resposta"] == "conforme" else p["nc"]

        mapeadas = {m[1]["id"] for m in mapa.values()}
        extras = {}
        for r in rs:
            if r["pergunta_id"] not in mapeadas:
                extras[r["pergunta_id"]] = "Conforme" if r["resposta"] == "conforme" else p["nc"]
                extras_texto[aba].setdefault(r["pergunta_id"], r["pergunta"])

        # evidências (fotos do checklist) e demais colunas de apoio
        urls = []
        for f in sorted(fotos_sub.get(s["id"], []), key=lambda x: x["sort_order"]):
            k = f["r2_key"] or ""
            urls.append(k if k.startswith("http") else (f"{base_r2}/{k}" if base_r2 and k else ""))
        put("evidencias", "; ".join(u for u in urls if u) or None)
        put("todos as inconformidades", "Sim" if all(r["resolvido"] is True for r in ncs) else "Não")
        cats = []
        for r in ncs:
            if r["categoria"] and r["categoria"] not in cats:
                cats.append(r["categoria"])
        put("categoria inconforme", ", ".join(cats) or None)
        obs = [r["observacao"].strip() for r in ncs if r.get("observacao") and r["observacao"].strip()]
        put("inconformidade da categoria", " | ".join(obs) or None)
        if aba == "Operacional":
            put("pergunta", " | ".join(r["pergunta"] for r in ncs) or None)

        d = local.day
        f1, f2, f3 = semanas.get((local.year, local.month), (8, 15, 22))
        put("semana", SEMANAS[0 if d <= f1 else 1 if d <= f2 else 2 if d <= f3 else 3])
        put("mes", dt.datetime(local.year, local.month, 1))
        # Na aba Operacional atual, "Gerente da Equipe" / "Gerente Obs" guardam o nome do COORDENADOR (como no BI).
        put("coordenador da equipe", hh.get("coordenador") or None)
        put("gerente da equipe", hh.get("coordenador") or None)
        put("base da equipe", hh.get("base") or s["base"])
        put("observador", s["observador"])
        emp = empregados.get(s["matricula"], {})
        put("coordenador do obs", emp.get("coordenador") or None)
        put("gerente obs", emp.get("coordenador") or None)
        put("base do obs", emp.get("base") or None)

        # colunas calculadas do final (valor fixo por tipo; vai na coluna da categoria OU na de inconformidade)
        const = p["const"].get(ger_txt, {})
        cats_nc = {categoria_chave(r["categoria"]) for r in ncs}
        i_res = p["i_resumo"]
        i = i_res
        while i < len(cols):
            nome = str(cols[i]).strip()
            nn = n(nome)
            if nn.startswith("inconformidade"):
                i += 1
                continue
            if nn in ("conformidade", "incondormidade", "indice de conformidade", "total", "indice"):
                break
            c = const.get(i, 0)
            tem_nc = categoria_chave(nome) in cats_nc
            row[i] = 0 if tem_nc else c
            if i + 1 < len(cols) and n(cols[i + 1]).startswith("inconformidade"):
                row[i + 1] = c if tem_nc else 0
            i += 1
        total = next((const.get(j, 0) for j, c in enumerate(cols) if n(c).strip() == "total"), 0)
        for j, c in enumerate(cols):
            nn = n(c).strip()
            if nn == "conformidade":
                row[j] = 0 if ncs else total
            elif nn == "incondormidade":
                row[j] = total if ncs else 0
            elif nn == "indice de conformidade":
                row[j] = 0 if ncs else 1
            elif nn == "total":
                row[j] = total
            elif nn == "indice":
                contador[aba] += 1
                row[j] = contador[aba]

        linhas[aba].append((row, extras, VALIDACAO.get(s["status"], s["status"])))

    # ── escrita ──
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for aba in ("Operacional", "Administrativo"):
        ws = wb.create_sheet(aba)
        ids_extras = sorted(extras_texto[aba])
        ws.append([str(c) for c in prep[aba]["cols"]] + [extras_texto[aba][i] for i in ids_extras] + ["Validação"])
        for row, extras, validacao in linhas[aba]:
            ws.append(row + [extras.get(i) for i in ids_extras] + [validacao])
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws["A1"].number_format = "General"
        for rr in range(2, ws.max_row + 1):
            ws.cell(rr, 1).number_format = "dd/mm/yyyy hh:mm:ss"
        i_mes = next(i for i, c in enumerate(prep[aba]["cols"], start=1) if n(c).strip() == "mes")
        for rr in range(2, ws.max_row + 1):
            ws.cell(rr, i_mes).number_format = "dd/mm/yyyy"
    wb.save(args.saida)
    print({a: len(v) for a, v in linhas.items()}, "| sem tipo conhecido:", dict(sem_tipo))
    print("perguntas extras (fora do modelo):", {a: len(v) for a, v in extras_texto.items()})
    print("arquivo:", args.saida)


if __name__ == "__main__":
    main()
