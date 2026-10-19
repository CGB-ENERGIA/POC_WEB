"""Exporta a Matriz de Responsabilidade de um mês no mesmo modelo da planilha "Matriz de Responsabilidade.xlsx".

Mesmas regras da tela de Matriz do painel: só checklists APROVADOS do mês, só não conformidades que não foram
resolvidas no ato; uma linha por checklist. "Resolvido" = Sim quando todas essas NCs têm resolução aprovada.

Uso: python exportar_matriz.py --ano 2026 --mes 10 --saida "C:\\caminho\\arquivo.xlsx"
Só lê o banco (chave anon do POC_OP/.env).
"""
import argparse
import datetime as dt
import os
import sys

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import exportar_mes as em  # noqa: E402  (paginar, em_lotes, FUSO)
import importar_historico as h  # noqa: E402

CABECALHO = ["Data", "Base", "Observador", "Equipe", "Categoria", "Inconformidade da categoria",
             "Status", "Resolvido", "Data de resolução", "Evidência"]
LARGURAS = [18, 9.8, 15.8, 14.7, 37.3, 48.2, 8.7, 14, 21.3, 24]


def local(iso):
    return dt.datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(em.FUSO).replace(tzinfo=None, microsecond=0)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ano", type=int, required=True)
    ap.add_argument("--mes", type=int, required=True)
    ap.add_argument("--saida", required=True)
    args = ap.parse_args()

    key = h.chave_env()
    ini = dt.datetime(args.ano, args.mes, 1, tzinfo=em.FUSO)
    fim = dt.datetime(args.ano + (args.mes == 12), args.mes % 12 + 1, 1, tzinfo=em.FUSO)
    fmt = lambda d: d.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    subs = em.paginar("checklist_submissions", key, {
        "select": "id,observador,data,base,equipe",
        "status": "eq.aprovado",
        "and": f"(data.gte.{fmt(ini)},data.lt.{fmt(fim)})",
    }, "data.asc,id.asc")
    ids = [s["id"] for s in subs]
    ncs = em.em_lotes("checklist_responses", key, "submission_id", ids,
                      "submission_id,pergunta_id,categoria,pergunta,observacao,resolvido",
                      extra={"resposta": "eq.nao_conforme"}, ordem="submission_id.asc,pergunta_id.asc")
    resolucoes = em.em_lotes("nc_resolucoes", key, "submission_id", ids,
                             "submission_id,pergunta_id,status,data_resolucao,observacao,comentario_analise", ordem="id.asc")
    res = {(r["submission_id"], r["pergunta_id"]): r for r in resolucoes}
    por_sub = {}
    for r in ncs:
        if r["resolvido"] is True:  # resolvida no ato pelo auditor: não entra na matriz
            continue
        por_sub.setdefault(r["submission_id"], []).append(r)

    linhas = []
    for s in subs:
        lista = por_sub.get(s["id"])
        if not lista:
            continue
        cats, textos = [], []
        for r in lista:
            if r["categoria"] and r["categoria"] not in cats:
                cats.append(r["categoria"])
            t = (r["observacao"] or "").strip() or r["pergunta"]
            if t not in textos:
                textos.append(t)
        rs = [res.get((s["id"], r["pergunta_id"])) for r in lista]
        resolvido = all(x is not None and x["status"] == "aprovado" for x in rs)
        data_res = evid = None
        if resolvido:
            data_res = max(local(x["data_resolucao"]) for x in rs).replace(hour=0, minute=0, second=0, microsecond=0)
            notas = []
            for x in rs:
                n = (x["comentario_analise"] or x["observacao"] or "").strip()
                if n and n not in notas:
                    notas.append(n)
            evid = " | ".join(notas) or None
        linhas.append([local(s["data"]), s["base"], s["observador"], s["equipe"], ", ".join(cats), " | ".join(textos),
                       "Não", "Sim" if resolvido else "Não", data_res, evid])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Página1"
    ws.append(CABECALHO)
    for l in linhas:
        ws.append(l)
    for c in ws[1]:
        c.font = Font(bold=True, color="595959")
    for i, w in enumerate(LARGURAS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for rr in range(2, ws.max_row + 1):
        ws.cell(rr, 1).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(rr, 9).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    wb.save(args.saida)
    print(f"{len(subs)} checklists aprovados | {len(linhas)} na matriz | resolvidos: {sum(1 for l in linhas if l[7] == 'Sim')} | arquivo: {args.saida}")


if __name__ == "__main__":
    main()
